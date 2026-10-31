"""
Word 导出

以 template.docx 为底稿（不修改原件），套入节正文，覆盖写入 export.docx。
规则见 docs/document-model.md。
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from deepdocument.document.model import Brief, Literature
from deepdocument.document.store import load_literature, load_meta, load_section, project_dir

EXPORTABLE_STATUSES = frozenset({"outlined", "draft", "accepted"})


def export_docx(project_slug: str, brief: Brief | None = None) -> Path:
    """
    导出 Word，返回 export.docx 路径

    Args:
        project_slug: 项目 slug
        brief: 可选；为空则从磁盘读取
    """
    from deepdocument.document.store import load_brief

    if brief is None:
        brief = load_brief(project_slug)
    if brief is None:
        raise FileNotFoundError(f"项目缺少 brief.md: {project_slug}")

    root = project_dir(project_slug)
    template_path = root / "template.docx"
    export_path = root / "export.docx"
    export_path.parent.mkdir(parents=True, exist_ok=True)

    if template_path.exists():
        doc = Document(str(template_path))
        # 清空示例正文，保留样式
        body = doc.element.body
        for child in list(body):
            body.remove(child)
    else:
        doc = Document()

    heading1 = _find_style(doc, ["Heading 1", "标题 1"])
    heading2 = _find_style(doc, ["Heading 2", "标题 2"])
    normal = _find_style(doc, ["Normal", "正文"])

    title_para = doc.add_heading(brief.title, level=1)
    if heading1:
        title_para.style = heading1

    meta = load_meta(project_slug)
    for section in meta.get("sections", []):
        status = section.get("status", "empty")
        if status not in EXPORTABLE_STATUSES:
            continue

        heading_text = section.get("title", "")
        if status == "draft":
            heading_text = f"{heading_text}（草稿）"
        elif status == "outlined":
            heading_text = f"{heading_text}（待写）"

        h = doc.add_heading(heading_text, level=2)
        if heading2:
            h.style = heading2

        content = load_section(project_slug, section["id"]).strip()
        if not content:
            continue
        for para_text in content.split("\n\n"):
            text = para_text.strip()
            if not text:
                continue
            p = doc.add_paragraph(text)
            if normal:
                p.style = normal

    _add_references(doc, project_slug, brief.citation_style, heading1)
    doc.save(str(export_path))
    return export_path


def _find_style(doc: Document, names: list[str]) -> str | None:
    styles = {s.name for s in doc.styles if s.type == WD_STYLE_TYPE.PARAGRAPH}
    for name in names:
        if name in styles:
            return name
    return None


def _add_references(
    doc: Document,
    project_slug: str,
    citation_style: str,
    heading_style: str | None,
) -> None:
    accepted = load_literature(project_slug, status="accepted")
    if not accepted:
        return

    h = doc.add_heading("参考文献", level=1)
    if heading_style:
        h.style = heading_style

    accepted = sorted(
        accepted,
        key=lambda lit: (lit.authors[0] if lit.authors else "", lit.year or 0),
    )
    for lit in accepted:
        doc.add_paragraph(_format_citation(lit, citation_style))


def _format_citation(lit: Literature, style: str) -> str:
    """第一版简单拼接；APA / GB 细规则后续再补"""
    authors = lit.authors or ["佚名"]
    if len(authors) <= 3:
        authors_str = ", ".join(authors)
    else:
        authors_str = ", ".join(authors[:3]) + ", et al."

    year = lit.year if lit.year is not None else "n.d."
    parts = [f"{authors_str}. ({year}). {lit.title}."]
    if lit.venue:
        parts.append(f" {lit.venue}.")
    if lit.doi:
        parts.append(f" https://doi.org/{lit.doi}")
    elif lit.url:
        parts.append(f" {lit.url}")

    # style 参数留给后续分支；当前输出对 APA / 简易 GB 都可读
    _ = style
    return "".join(parts)
